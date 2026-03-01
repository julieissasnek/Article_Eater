#!/usr/bin/env python3
"""
Generate a PPTX presentation with all website content choices
for David Kirsh's personal website (kirsh.anfarch.org).

Reads from data/website/*.json and produces a slide deck
that can be reviewed page-by-page in PowerPoint/Keynote.
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '.pptx_lib'))

import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DATA = Path(__file__).parent.parent / "data" / "website"
OUT = DATA / "MY_Website_Content_Review.pptx"

# ── Colors ──
DARK_BG  = RGBColor(0x1A, 0x1A, 0x2E)
ORANGE   = RGBColor(0xE8, 0x6C, 0x37)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT    = RGBColor(0xCC, 0xCC, 0xCC)
GREEN    = RGBColor(0x4E, 0xC9, 0xB0)
RED      = RGBColor(0xE0, 0x4E, 0x4E)
YELLOW   = RGBColor(0xFF, 0xD7, 0x00)


def set_slide_bg(slide, color=DARK_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, left, top, width, height,
             font_size=14, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf


def add_bullet(tf, text, font_size=12, color=LIGHT, indent=0):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.level = indent
    return p


def title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide)
    add_text(slide, title, 0.5, 2.5, 9, 1.5, font_size=36, bold=True,
             color=ORANGE, alignment=PP_ALIGN.CENTER)
    if subtitle:
        add_text(slide, subtitle, 0.5, 4.2, 9, 1, font_size=16,
                 color=LIGHT, alignment=PP_ALIGN.CENTER)
    return slide


def section_header(prs, section_name, status_emoji, status_text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text(slide, f"{status_emoji}  {section_name}", 0.5, 0.3, 9, 0.8,
             font_size=28, bold=True, color=ORANGE)
    add_text(slide, f"Status: {status_text}", 0.5, 1.1, 9, 0.5,
             font_size=14, color=LIGHT)
    return slide


def research_topic_slide(prs, topic):
    # Slide 1: Overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    add_text(slide, f"RESEARCH: {topic['title']}", 0.5, 0.2, 9, 0.6,
             font_size=24, bold=True, color=ORANGE)

    # Short abstract
    add_text(slide, "Card Preview (short abstract):", 0.5, 0.9, 9, 0.4,
             font_size=10, color=GREEN, bold=True)
    add_text(slide, topic['short_abstract'], 0.5, 1.3, 9, 1.5,
             font_size=12, color=WHITE)

    # Key findings
    y = 3.0
    tf = add_text(slide, "Key Findings:", 0.5, y, 9, 0.3,
                  font_size=10, color=GREEN, bold=True)
    for finding in topic['key_findings']:
        add_bullet(tf, f"• {finding}", font_size=11, color=LIGHT)

    # Related + video
    related = ", ".join(topic.get('related_topics', []))
    video = topic.get('video_placeholder', 'No video specified')
    add_text(slide, f"Related: {related}", 0.5, 6.5, 4.5, 0.3,
             font_size=9, color=LIGHT)
    add_text(slide, f"🎬 {video}", 5.0, 6.5, 4.5, 0.3,
             font_size=9, color=YELLOW)

    # Slide 2: Long summary
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide2)
    add_text(slide2, f"RESEARCH: {topic['title']} — Full Page Text", 0.5, 0.2, 9, 0.6,
             font_size=20, bold=True, color=ORANGE)

    # Split long summary into paragraphs
    paragraphs = topic['long_summary'].split('\n\n')
    y = 1.0
    for para in paragraphs:
        if y > 6.5:
            break
        height = max(0.8, len(para) / 200)
        add_text(slide2, para, 0.5, y, 9, height,
                 font_size=11, color=WHITE)
        y += height + 0.15

    return slide, slide2


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ═══════════════════════════════════
    # TITLE SLIDE
    # ═══════════════════════════════════
    title_slide(prs,
                "MY Website — Content Review",
                "kirsh.anfarch.org  •  David Kirsh  •  2026-02-25\n"
                "Review each slide to approve, edit, or flag content for revision.")

    # ═══════════════════════════════════
    # TABLE OF CONTENTS
    # ═══════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text(slide, "Content Plan Overview", 0.5, 0.3, 9, 0.6,
             font_size=24, bold=True, color=ORANGE)

    toc_items = [
        ("🔴", "Homepage", "Partial — needs Travel Schedule, hero section update"),
        ("🔴", "Research: Architecture", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Choreography", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Creativity", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Embodied/Extended/Situated", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Interactivity", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Planning", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Representation", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Perceptual Thinking", "Lorem Ipsum → NEW content generated"),
        ("🔴", "Research: Implicit/Explicit", "Lorem Ipsum → NEW content generated"),
        ("✅", "Lab", "Real content — could expand with team/projects"),
        ("✅", "Publications", "Real content — comprehensive, categorized"),
        ("🟡", "Teaching", "Stale (Fall '21) — needs current courses"),
        ("🔴", "Speaker/Consultant", "Entirely Lorem Ipsum + placeholder video"),
        ("✅", "About Me", "Rich narrative — complete"),
        ("🆕", "CMR / Web of Belief", "New section from Article Eater repo"),
    ]

    tf = add_text(slide, "", 0.5, 1.1, 9, 5.5, font_size=11, color=WHITE)
    tf.paragraphs[0].text = ""
    for emoji, name, status in toc_items:
        add_bullet(tf, f"{emoji}  {name:40s} {status}", font_size=11, color=LIGHT)

    # ═══════════════════════════════════
    # HOMEPAGE
    # ═══════════════════════════════════
    slide = section_header(prs, "Homepage", "🏠", "Partial — hero works, Travel Schedule empty")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Current content (keep):"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GREEN
    add_bullet(tf, "Hero: Rothko-inspired sunset image with profile photo ✅")
    add_bullet(tf, '"David Kirsh — PROFESSOR • RESEARCHER • CONSULTANT" ✅')
    add_bullet(tf, "Who Am I: Professor UCSD, President ANFA, Visiting Prof UCL Bartlett ✅")
    add_bullet(tf, "Research Interests: Embodied cognition, Design thinking, Visual thinking, Creativity ✅")
    add_bullet(tf, "Education: MIT AI Lab, Oxford D.Phil, U of Toronto ✅")
    add_bullet(tf, "")
    p = add_bullet(tf, "Needs updating:", font_size=12, color=RED)
    p.font.bold = True
    add_bullet(tf, '🔴 "Currently Working On" → Update to CMR / Web of Belief project')
    add_bullet(tf, "🔴 Travel Schedule dropdown → Empty, needs 2026 dates or removal")
    add_bullet(tf, "🔴 Research Interest buttons → Link to research sub-pages (functional but verify)")
    add_bullet(tf, "")
    p = add_bullet(tf, "🎬 Suggested video: None on homepage (keep clean)", color=YELLOW)

    # ═══════════════════════════════════
    # 9 RESEARCH TOPICS
    # ═══════════════════════════════════
    research = json.loads((DATA / "research_content.json").read_text())
    for topic in research['topics']:
        research_topic_slide(prs, topic)

    # ═══════════════════════════════════
    # LAB
    # ═══════════════════════════════════
    slide = section_header(prs, "Lab — Interactive Cognition Lab", "🔬", "Complete — could expand")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Current content (keep):"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GREEN
    add_bullet(tf, "Lab name: Interactive Cognition Lab ✅")
    add_bullet(tf, "Description of focus: cognitive science experiments, computational models, interactive installations ✅")
    add_bullet(tf, "Address: SSRB 100, 9500 Gilman Drive, UCSD ✅")
    add_bullet(tf, "Google Map embed: working ✅")
    add_bullet(tf, "Contact info: kirsh@ucsd.edu ✅")
    add_bullet(tf, "")
    p = add_bullet(tf, "Could add:", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🟡 Team members / current students")
    add_bullet(tf, "🟡 Current projects list (Article Eater / CMR, VR experiments)")
    add_bullet(tf, "🟡 Lab photos (155 images available in ~/Documents/Lab/)")
    add_bullet(tf, "🆕 Link to QA Browse System / Knowledge Base")

    # ═══════════════════════════════════
    # PUBLICATIONS
    # ═══════════════════════════════════
    slide = section_header(prs, "Publications", "📚", "Complete — real categorized list")
    cv = json.loads((DATA / "cv_parsed.json").read_text())
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = f"Current: {len(cv.get('publications', []))} publications parsed from CV"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GREEN
    add_bullet(tf, "Website already has real publication list with Abstract + PDF links (Google Drive) ✅")
    add_bullet(tf, "Topics covered: Architecture, Choreography, Creativity, Embodied, Interactivity, etc. ✅")
    add_bullet(tf, "")
    p = add_bullet(tf, "Could improve:", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🟡 Cross-link publications to their research topic pages")
    add_bullet(tf, "🟡 Add citation counts (from Google Scholar)")
    add_bullet(tf, "🟡 Ensure PDF links are still active")
    add_bullet(tf, "🟡 Add post-2021 publications from CV")

    # Sample pubs
    add_bullet(tf, "")
    add_bullet(tf, "Sample publications from CV:", color=GREEN)
    for pub in cv.get('publications', [])[:5]:
        year = pub.get('year', '?')
        title = pub.get('title', '')[:80]
        add_bullet(tf, f"  [{year}] {title}", font_size=10, color=LIGHT)

    # ═══════════════════════════════════
    # TEACHING
    # ═══════════════════════════════════
    slide = section_header(prs, "Teaching", "🎓", "Stale — showing Fall '21 courses")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Current content (stale):"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RED
    add_bullet(tf, "COGS 100: Cyborgs Now and in the Future")
    add_bullet(tf, "COGS 187A/B: Usability & Information Architecture")
    add_bullet(tf, "COGS 199: Independent Study")
    add_bullet(tf, "COGS 260: 4E Cognition")
    add_bullet(tf, "Canvas links shown but may be outdated")
    add_bullet(tf, "")
    p = add_bullet(tf, "Action required:", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🔴 USER TO PROVIDE: Current semester courses and times")
    add_bullet(tf, "🟡 Add course descriptions (23 course folders found in ~/Documents)")
    add_bullet(tf, "🟡 Update Canvas links")

    # ═══════════════════════════════════
    # SPEAKER/CONSULTANT
    # ═══════════════════════════════════
    slide = section_header(prs, "Speaker / Consultant", "🎤", "Entirely placeholder — needs everything")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Current content: ALL Lorem Ipsum + placeholder React tutorial video"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RED
    add_bullet(tf, "")
    p = add_bullet(tf, "Proposed content:", font_size=12, color=GREEN)
    p.font.bold = True
    add_bullet(tf, "Speaking topics: 4E Cognition, Architecture & the Brain, Creativity, Dance & Thinking")
    add_bullet(tf, f"From CV: 61 grants, 8 major awards document consulting expertise")
    add_bullet(tf, "Key affiliations: ANFA President, UCL Bartlett Leverhulme Fellow")
    add_bullet(tf, "Notable clients/orgs: Steelcase, Rockefeller Foundation, IBM, DARPA")
    add_bullet(tf, "")
    p = add_bullet(tf, "🎬 VIDEOS NEEDED (user to provide YouTube links):", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🎬 Leverhulme Lecture Series (Bartlett UCL, 2017–2019)")
    add_bullet(tf, "🎬 V&A Museum 'Thinking with Tools' keynote (2019)")
    add_bullet(tf, "🎬 Any TEDx or public lecture recording")
    add_bullet(tf, "🎬 ANFA conference keynotes")

    # ═══════════════════════════════════
    # ABOUT ME
    # ═══════════════════════════════════
    slide = section_header(prs, "About Me", "👤", "Complete — rich personal narrative")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Current content (keep):"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GREEN
    add_bullet(tf, "Personally: Early interest in Ouspenski/Gurdjieff, shift from economics to philosophy ✅")
    add_bullet(tf, "Artist's Entrance: Reflections on making life a work of art ✅")
    add_bullet(tf, "Traveling: Meditation and breaking sameness ✅")
    add_bullet(tf, "Food Delight: Personal interests ✅")
    add_bullet(tf, "Working for Money vs. Fun: Essay ✅")
    add_bullet(tf, "3D Video Game: Thought experiment ✅")
    add_bullet(tf, "")
    p = add_bullet(tf, "Could add:", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🟡 Updated biography reflecting current work (CMR, ANFA presidency)")
    add_bullet(tf, "🟡 Professional headshot (current one works but could be updated)")

    # ═══════════════════════════════════
    # NEW: CMR / WEB OF BELIEF
    # ═══════════════════════════════════
    slide = section_header(prs, "NEW SECTION: CMR / Web of Belief", "🆕",
                           "New — showcase the Article Eater project")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=12, color=WHITE)
    tf.paragraphs[0].text = "Proposed new section:"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GREEN
    add_bullet(tf, "A Web of Belief for Cognitive & Neuroscience for Architecture")
    add_bullet(tf, "186 evidence-backed design templates, 13 thematic molecules, 812+ papers")
    add_bullet(tf, "Interactive Knowledge Base with search and topic pages")
    add_bullet(tf, "MASTER_DOC: 8,662 lines of original prose on how neuroscience informs design")
    add_bullet(tf, "")
    p = add_bullet(tf, "Content available:", font_size=12, color=GREEN)
    p.font.bold = True
    add_bullet(tf, "✅ Project overview from MASTER_DOC §1–2")
    add_bullet(tf, "✅ 30 worked examples (e.g., Salk Institute, Fallingwater, Guggenheim)")
    add_bullet(tf, "✅ Interactive QA Browse System (could embed or link)")
    add_bullet(tf, "✅ Citation graph visualization (1,735 edges)")
    add_bullet(tf, "")
    p = add_bullet(tf, "🎬 Videos:", font_size=12, color=YELLOW)
    p.font.bold = True
    add_bullet(tf, "🎬 ANFA talks on neuroscience for architecture")
    add_bullet(tf, "🎬 Bartlett lectures on evidence-based design")

    # ═══════════════════════════════════
    # VIDEO INVENTORY
    # ═══════════════════════════════════
    slide = section_header(prs, "Video Strategy — Talks to Embed", "🎬",
                           "User to provide YouTube links")
    tf = add_text(slide, "", 0.5, 1.8, 9, 5, font_size=11, color=WHITE)
    tf.paragraphs[0].text = "Videos needed for each research page:"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = YELLOW

    video_map = [
        ("Architecture", "Leverhulme/Bartlett lecture on architecture + cognition; ANFA talk"),
        ("Choreography", "Wayne McGregor collaboration video; Wellcome Foundation talk; V&A keynote"),
        ("Creativity", "Florence 2018 creativity talk; Viasat study presentation; Laban lecture"),
        ("Embodied/Extended", "CHI 2011 embodied cognition workshop; Dagstuhl on-body interaction talk"),
        ("Interactivity", "Theory of Interactivity lecture (Nice 2019 or UCL Bartlett)"),
        ("Planning", "MIT AI Lab era talk (if recorded) or more recent talk on planning"),
        ("Representation", "Talk on visual reasoning, diagrams, or notation design"),
        ("Perceptual Thinking", "Talk on imagination, projection, or design cognition"),
        ("Implicit/Explicit", "Talk on implicit cognition or skilled performance"),
        ("Speaker page", "Curated highlight reel or best single talk (keynote quality)"),
        ("CMR / Web of Belief", "ANFA neuroscience for architecture talk"),
    ]

    for page, desc in video_map:
        add_bullet(tf, f"{page}: {desc}", font_size=10, color=LIGHT)

    add_bullet(tf, "")
    add_bullet(tf, "Also check: ~/Documents/DK Website/ Video's/ for existing assets", color=YELLOW)
    add_bullet(tf, "Also check: YouTube channel (if any) for existing uploads", color=YELLOW)

    # ═══════════════════════════════════
    # FINAL SLIDE
    # ═══════════════════════════════════
    title_slide(prs,
                "Next Steps",
                "1. Review these slides — mark up anything to change\n"
                "2. Provide the React codebase (zip)\n"
                "3. Provide YouTube links for talk videos\n"
                "4. Confirm current courses for Teaching page\n"
                "5. We inject content, fix images, deploy")

    # Save
    prs.save(str(OUT))
    print(f"✅ Saved {OUT}")
    print(f"   {len(prs.slides)} slides total")


if __name__ == "__main__":
    main()
